#!/usr/bin/env python3
"""Test suite for the academic-poster scripts (stdlib unittest, Python >= 3.8).

Run:  python3 tests/make_fixtures.py && python3 tests/run_tests.py
Fixtures are synthetic; every expected value is derived by hand in tests/make_fixtures.py and
tests/README.md, never copied from script output. Assertions parse the JSON the scripts print and
compare numbers with an explicit tolerance:
  PPI_TOL = 0.1   (PPI is rounded to 0.1 by the scripts)
  PT_TOL  = 0.1   (printed pt rounded to 0.1)
  MM_TOL  = 0.6   (margins are measured on a 50-dpi render: 0.51 mm pixels + anti-aliasing)
  RATIO_TOL = 0.01 (contrast ratio rounded to 0.01)
Tests whose optional dependency (PyMuPDF, python-pptx, OpenCV/pyzbar, Playwright) is missing SKIP.
"""
import ast
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SKILL = os.path.join(ROOT, "skills", "academic-poster")
SCRIPTS = os.path.join(SKILL, "scripts")
TEMPLATES = os.path.join(SKILL, "templates")
FX = os.path.join(HERE, "fixtures")

PPI_TOL, PT_TOL, MM_TOL, RATIO_TOL = 0.1, 0.1, 0.6, 0.01
MM = 72 / 25.4
QR_GOOD = "https://example.org/poster"


def has(mod):
    try:
        return importlib.util.find_spec(mod) is not None
    except (ImportError, ValueError):
        return False


HAS_FITZ, HAS_PPTX = has("pymupdf") or has("fitz"), has("pptx")
HAS_QR_DECODER = (has("cv2") and has("numpy")) or (has("pyzbar") and has("PIL"))


def manifest():
    with open(os.path.join(FX, "manifest.json"), encoding="utf-8") as fh:
        return json.load(fh)


def run(script, *args, **kw):
    """Run a script; return (exit code, parsed JSON or None, raw stdout)."""
    env = dict(os.environ)
    env.update(kw.get("env", {}))
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + list(args),
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env, cwd=HERE)
    out = p.stdout.decode("utf-8", "replace")
    try:
        data = json.loads(out)
    except ValueError:
        data = None
    return p.returncode, data, out


def fx(name):
    return os.path.join(FX, name)


def check(data, cid):
    for c in data["checks"]:
        if c["id"] == cid:
            return c
    raise AssertionError("check %r not reported" % cid)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="apo_test_")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


# --------------------------------------------------------------------------- every script
class TestHelp(unittest.TestCase):
    def test_help_exits_zero(self):
        for s in sorted(os.listdir(SCRIPTS)):
            if s.endswith(".py"):
                code, _, out = run(s, "--help")
                self.assertEqual(code, 0, s)
                self.assertIn("usage", out.lower(), s)


# --------------------------------------------------------------------------- verify_wording
class TestVerifyWording(Base):
    def test_pass_case(self):
        code, d, _ = run("verify_wording.py", fx("wording_pass.json"), "--source", fx("paper.tex"))
        self.assertEqual(code, 0)
        self.assertIs(d["ok"], True)
        self.assertEqual((d["fragments_failing"], d["short_failing"], d["source_lines"]), (0, 0, 10))

    def test_fail_case_counts(self):
        # 1 order failure; 1 over-limit short (7 > 5); 1 bad-line short (99 > 10 lines)
        code, d, _ = run("verify_wording.py", fx("wording_fail.json"), "--source", fx("paper.tex"))
        self.assertEqual(code, 1)
        self.assertEqual((d["fragments_failing"], d["short_failing"]), (1, 2))
        r = {x["slot"]: x for x in d["results"]}
        self.assertTrue(r["ok_slot"]["fragments_ok"])
        self.assertEqual(r["order_slot"]["missing"][0]["missing_word"], "models")
        self.assertEqual((r["long_short"]["short_words"], r["long_short"]["limit"]), (7, 5))
        self.assertEqual(r["badline_short"]["bad_lines"], [99])
        self.assertEqual(r["badline_short"]["limit"], 20)          # --default-limit

    def test_free_wording(self):
        # fragments skipped; only the word cap 7 > 5 still fails
        code, d, _ = run("verify_wording.py", fx("wording_fail.json"), "--source", fx("paper.tex"), "--free-wording")
        self.assertEqual(code, 1)
        self.assertEqual((d["fragments_failing"], d["short_failing"], d["mode"]), (0, 1, "free-wording"))

    def test_window_zero_still_finds_same_line(self):
        code, d, _ = run("verify_wording.py", fx("wording_pass.json"), "--source", fx("paper.tex"), "--window", "0")
        self.assertEqual(code, 0)

    def test_limits_file(self):
        # precedence: slot "limit" > --limits file > --default-limit.
        # caption (3 words) has its own limit 5, so a --limits cap of 1 must not apply -> pass
        lim = os.path.join(self.tmp, "lim.json")
        with open(lim, "w") as fh:
            json.dump({"caption": 1, "badline_short": 2}, fh)
        code, d, _ = run("verify_wording.py", fx("wording_pass.json"), "--source", fx("paper.tex"), "--limits", lim)
        self.assertEqual(code, 0)
        # badline_short has no own limit -> the --limits value 2 applies
        code, d, _ = run("verify_wording.py", fx("wording_fail.json"), "--source", fx("paper.tex"), "--limits", lim)
        r = {x["slot"]: x for x in d["results"]}
        self.assertEqual(r["badline_short"]["limit"], 2)


# --------------------------------------------------------------------------- check_claims
class TestCheckClaims(unittest.TestCase):
    def test_unregistered_number_fails(self):
        code, d, _ = run("check_claims.py", "--poster", fx("poster_text.txt"), "--source", fx("paper.tex"))
        self.assertEqual(code, 1)
        self.assertEqual((d["numbers_checked"], d["failing"]), (6, 1))
        status = {c["normalised"]: c["status"] for c in d["checks"]}
        self.assertEqual(status, {"0.64": "PASS", "143.5%": "PASS", "1250": "PASS", "25": "PASS",
                                  "2500": "PASS", "42%": "FAIL"})

    def test_registered_with_derivation_passes(self):
        code, d, _ = run("check_claims.py", "--poster", fx("poster_text.txt"), "--source", fx("paper.tex"),
                         "--registered", fx("claims_registered.json"))
        self.assertEqual(code, 0)
        self.assertEqual((d["numbers_checked"], d["failing"]), (6, 0))

    def test_registered_without_derivation_still_fails(self):
        code, d, _ = run("check_claims.py", "--poster", fx("poster_text.txt"), "--source", fx("paper.tex"),
                         "--registered", fx("claims_noderiv.json"))
        self.assertEqual((code, d["failing"]), (1, 1))

    def test_ignore(self):
        code, d, _ = run("check_claims.py", "--poster", fx("poster_text.txt"), "--source", fx("paper.tex"),
                         "--ignore", "42%")
        self.assertEqual((code, d["numbers_checked"]), (0, 5))

    def test_wording_json_route(self):
        code, d, _ = run("check_claims.py", "--poster", fx("claims_wording.json"), "--source", fx("paper.tex"))
        self.assertEqual((code, d["numbers_checked"], d["failing"]), (0, 2, 0))

    @unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
    def test_pdf_poster_route(self):
        # clean.pdf text: "Clean poster title" / "Body text at thirty points" -> no numbers at all
        code, d, _ = run("check_claims.py", "--poster", fx("clean.pdf"), "--source", fx("paper.tex"))
        self.assertEqual((code, d["numbers_checked"]), (0, 0))


# --------------------------------------------------------------------------- export_figs
class TestExportFigs(Base):
    def test_raster_stdlib(self):
        code, d, _ = run("export_figs.py", fx("fig_raster.png"), "--out", self.tmp, "--placed-cm", "2.54")
        self.assertEqual(code, 0)
        f = d["figures"][0]
        self.assertEqual(f["px"], [300, 200])
        self.assertAlmostEqual(f["aspect_h_over_w"], 200 / 300.0, places=5)
        self.assertAlmostEqual(f["ppi_at_placed"], 300.0, delta=PPI_TOL)     # 300 px over 1 in

    @unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
    def test_vector_pdf_long_side(self):
        # 288 x 216 pt -> 6000 x 4500 px; 6000 px over 25.4 cm (10 in) -> 600 PPI
        code, d, _ = run("export_figs.py", fx("fig_4x3.pdf"), "--out", self.tmp, "--placed-cm", "25.4", "--svg")
        self.assertEqual(code, 0)
        f = d["figures"][0]
        self.assertEqual(f["px"], [6000, 4500])
        self.assertAlmostEqual(f["aspect_h_over_w"], 0.75, places=5)
        self.assertAlmostEqual(f["ppi_at_placed"], 600.0, delta=PPI_TOL)
        self.assertTrue(os.path.getsize(f["png"]) > 0 and os.path.getsize(f["svg"]) > 0)

    @unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
    def test_long_side_option(self):
        code, d, _ = run("export_figs.py", fx("fig_4x3.pdf"), "--out", self.tmp, "--long-side", "1000")
        self.assertEqual(d["figures"][0]["px"], [1000, 750])                # 1000 * 216/288

    def test_missing_file_is_error(self):
        code, d, _ = run("export_figs.py", os.path.join(self.tmp, "nope.png"), "--out", self.tmp)
        self.assertEqual((code, d["ok"], len(d["errors"])), (1, False, 1))


# --------------------------------------------------------------------------- check_export
@unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
class TestCheckExport(unittest.TestCase):
    def test_clean_poster_passes(self):
        code, d, _ = run("check_export.py", fx("clean.pdf"), "--size", "A0", "--contrast")
        self.assertEqual(code, 0)
        self.assertIs(d["ok"], True)
        self.assertAlmostEqual(d["scale"], 1.0, places=4)
        for cid in ("size", "images", "text", "fonts", "margins", "contrast"):
            self.assertEqual(check(d, cid)["status"], "PASS", cid)
        self.assertAlmostEqual(check(d, "images")["detail"]["images"][0]["ppi"], 300.0, delta=PPI_TOL)
        self.assertAlmostEqual(check(d, "text")["detail"]["min_pt"], 30.0, delta=PT_TOL)
        self.assertAlmostEqual(check(d, "margins")["detail"]["ink_to_edge_mm"]["left"], 20.0, delta=MM_TOL)

    def test_size_formats_equivalent(self):
        # A0 == 84.1x118.9cm == 841x1189mm
        for size in ("84.1x118.9cm", "841x1189mm"):
            code, d, _ = run("check_export.py", fx("clean.pdf"), "--size", size)
            self.assertEqual(check(d, "size")["status"], "PASS", size)

    def test_same_aspect_is_warn_and_scaled(self):
        # A1 page checked as A0: scale 841/594 = 1.415825; 20 pt -> 28.32 pt
        code, d, _ = run("check_export.py", fx("a1_scaled.pdf"), "--size", "A0")
        self.assertEqual(code, 0)
        self.assertEqual(check(d, "size")["status"], "WARN")
        self.assertAlmostEqual(d["scale"], 841 / 594.0, places=4)
        self.assertAlmostEqual(check(d, "text")["detail"]["min_pt"], 28.3, delta=PT_TOL)

    def test_qr(self):
        code, d, _ = run("check_export.py", fx("clean.pdf"), "--size", "A0", "--qr-url", QR_GOOD)
        q = check(d, "qr")
        if not HAS_QR_DECODER:
            self.assertEqual(q["status"], "SKIPPED")      # degrade, never fail
            self.assertEqual(code, 0)
            return
        if not manifest().get("qr_generated"):
            self.skipTest("QR fixture not generated (OpenCV QRCodeEncoder unavailable)")
        self.assertEqual(q["status"], "PASS")
        self.assertIn(QR_GOOD, q["detail"]["decoded"])

    @unittest.skipUnless(HAS_QR_DECODER, "no QR decoder (opencv-python or pyzbar+Pillow)")
    def test_qr_absent_fails(self):
        # a1_scaled.pdf has no QR at all -> decoded [] -> FAIL
        code, d, _ = run("check_export.py", fx("a1_scaled.pdf"), "--size", "A0", "--qr-url", QR_GOOD)
        self.assertEqual((code, check(d, "qr")["status"], check(d, "qr")["detail"]["decoded"]), (1, "FAIL", []))

    def test_crops_written_for_failures(self):
        tmp = tempfile.mkdtemp(prefix="apo_crops_")
        try:
            run("check_export.py", fx("defects.pdf"), "--size", "A0", "--crops", tmp)
            names = sorted(os.listdir(tmp))
            self.assertTrue(any(n.startswith("img_") for n in names), names)
            self.assertTrue(any(n.startswith("span_") for n in names), names)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


@unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
class TestReviewDryRun(unittest.TestCase):
    """REVIEW mode, preflight lens: check_export on the planted-defect poster must report EVERY defect."""

    @classmethod
    def setUpClass(cls):
        cls.code, cls.d, _ = run("check_export.py", fx("defects.pdf"), "--size", "A0",
                                 "--qr-url", QR_GOOD, "--contrast")
        cls.s = 841 / (2400 / 72 * 25.4)          # 0.993307

    def test_overall_fail(self):
        self.assertEqual(self.code, 1)
        self.assertIs(self.d["ok"], False)
        self.assertAlmostEqual(self.d["scale"], self.s, places=4)

    def test_wrong_page_size(self):
        c = check(self.d, "size")
        self.assertEqual(c["status"], "FAIL")
        self.assertAlmostEqual(c["detail"]["delta_mm"][0], 2400 / 72 * 25.4 - 841, delta=0.01)  # +5.67
        self.assertAlmostEqual(c["detail"]["delta_mm"][1], 0.0, delta=0.01)

    def test_low_ppi_image(self):
        c = check(self.d, "images")
        self.assertEqual((c["status"], c["detail"]["below_min"]), ("FAIL", 1))
        self.assertAlmostEqual(c["detail"]["min_ppi"], 300 / (4 * self.s), delta=PPI_TOL)          # 75.5

    def test_tiny_text(self):
        c = check(self.d, "text")
        self.assertEqual((c["status"], len(c["detail"]["below_floor"])), ("FAIL", 1))
        self.assertAlmostEqual(c["detail"]["below_floor"][0]["printed_pt"], 12 * self.s, delta=PT_TOL)  # 11.9

    def test_unembedded_font(self):
        c = check(self.d, "fonts")
        self.assertEqual(c["status"], "FAIL")
        self.assertEqual([p["font"] for p in c["detail"]["problems"]], ["Helvetica"])

    def test_element_near_edge(self):
        c = check(self.d, "margins")
        self.assertEqual(c["status"], "FAIL")
        self.assertEqual(sorted(c["detail"]["too_close"]), ["left"])
        self.assertAlmostEqual(c["detail"]["too_close"]["left"], 5 * self.s, delta=MM_TOL)         # 4.97

    def test_low_contrast(self):
        c = check(self.d, "contrast")
        self.assertEqual((c["status"], len(c["detail"]["low_contrast"])), ("WARN", 1))
        self.assertAlmostEqual(c["detail"]["low_contrast"][0]["ratio"], 1.606, delta=RATIO_TOL)

    def test_wrong_qr(self):
        c = check(self.d, "qr")
        if not HAS_QR_DECODER:
            self.assertEqual(c["status"], "SKIPPED")
            return
        if not manifest().get("qr_generated"):
            self.skipTest("QR fixture not generated (OpenCV QRCodeEncoder unavailable)")
        self.assertEqual(c["status"], "FAIL")
        self.assertEqual(c["detail"]["decoded"], ["https://example.org/wrong"])


# --------------------------------------------------------------------------- build_pptx
@unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
class TestBuildPptx(Base):
    def test_example_spec(self):
        out = os.path.join(self.tmp, "example.pptx")
        code, d, _ = run("build_pptx.py", os.path.join(TEMPLATES, "plan-spec.example.json"), "--out", out,
                         "--wording", os.path.join(TEMPLATES, "wording.template.json"),
                         "--assets", fx("pptx_assets"))
        self.assertEqual(code, 0, d)
        self.assertEqual((d["problems"], d["elements"], d["page_cm"]), ([], 14, [84.1, 118.9]))
        from pptx import Presentation
        prs = Presentation(out)
        self.assertEqual((prs.slide_width, prs.slide_height), (int(84.1 * 360000), int(118.9 * 360000)))
        shapes = list(prs.slides[0].shapes)
        self.assertEqual(len(shapes), 14)
        # element 9 is intro_text, filled from the "intro" slot (two lines)
        self.assertEqual(shapes[9].text_frame.text,
                         "Small models are cheap to run\nbut they lose accuracy on rare classes")
        # results_caption is 34 px = 25.5 pt (>= 24)
        self.assertAlmostEqual(shapes[12].text_frame.paragraphs[0].runs[0].font.size.pt, 25.5, delta=0.01)
        # image: 1380 design px wide = 1380/96*2.54 = 36.5125 cm, h = w/2 (2760 x 1380 px)
        pic = shapes[5]
        self.assertAlmostEqual(pic.width / 360000.0, 1380 / 96.0 * 2.54, delta=0.01)
        self.assertAlmostEqual(pic.height / 360000.0, 1380 / 96.0 * 2.54 / 2, delta=0.01)

    def test_example_without_assets_reports_missing(self):
        # the shipped example points at figs/example_results.png, which is not shipped, and needs a wording file
        code, d, _ = run("build_pptx.py", os.path.join(TEMPLATES, "plan-spec.example.json"), "--check-only",
                         "--out", os.path.join(self.tmp, "x.pptx"))
        self.assertEqual(code, 1)
        kinds = sorted(p["problem"].split(":")[0] for p in d["problems"])
        self.assertEqual(kinds, ["file not found", "slot not found", "slot not found", "slot not found"])
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "x.pptx")))

    def test_bad_spec_problems(self):
        out = os.path.join(self.tmp, "bad.pptx")
        code, d, _ = run("build_pptx.py", fx("pptx_bad_spec.json"), "--out", out, "--assets", FX)
        self.assertEqual(code, 1)
        by = {p["element"]: p for p in d["problems"]}
        self.assertEqual(sorted(by), ["ghost", "low_img", "tiny", "wide_rect"])
        self.assertEqual(by["low_img"]["px"], [300, 150])
        self.assertEqual(by["low_img"]["placed_cm"], [10.0, 5.0])
        ppi = float(re.match(r"image (\d+) PPI", by["low_img"]["problem"]).group(1))
        self.assertAlmostEqual(ppi, 300 / (10 / 2.54), delta=0.6)       # 76.2, printed as integer
        self.assertTrue(os.path.exists(out))                            # still written


# --------------------------------------------------------------------------- render_script_text
class TestRenderScriptText(Base):
    def test_html_only_counts(self):
        out = os.path.join(self.tmp, "r.png")
        code, d, _ = run("render_script_text.py", fx("render_ok.json"), "--out", out, "--html-only")
        self.assertEqual(code, 0)
        self.assertEqual([(r["count"], r["ok"]) for r in d["rows"]], [(3, True), (3, True)])
        with open(d["html"], encoding="utf-8") as fh:
            html = fh.read()
        self.assertEqual(html.count("class='chip'"), 6)
        self.assertIn("สวัส", html)

    def test_count_assertion_fails(self):
        code, d, _ = run("render_script_text.py", fx("render_bad.json"), "--out",
                         os.path.join(self.tmp, "b.png"), "--html-only")
        self.assertEqual((code, d["ok"], d["rows"][0]["expected"]), (1, False, 4))

    @unittest.skipUnless(has("playwright"), "playwright not installed")
    def test_png_width_is_width_times_dpr(self):
        out = os.path.join(self.tmp, "r.png")
        code, d, raw = run("render_script_text.py", fx("render_ok.json"), "--out", out, "--dpr", "2")
        if code != 0 and d is None:
            self.skipTest("playwright present but Chromium could not launch")
        self.assertEqual(code, 0)
        self.assertEqual(d["px"][0], 400 * 2)
        import struct
        with open(out, "rb") as fh:
            w = struct.unpack(">II", fh.read(24)[16:24])[0]
        self.assertAlmostEqual(w, 800, delta=1)


# --------------------------------------------------------------------------- render_crops
HAS_PIL = has("PIL")
BOX = {"name": "intro", "x0": 0.1, "y0": 0.1, "x1": 0.6, "y1": 0.4}


def png_size(path):
    import struct
    with open(path, "rb") as fh:
        return list(struct.unpack(">II", fh.read(24)[16:24]))


class TestRenderCrops(Base):
    def crops_file(self, boxes):
        path = os.path.join(self.tmp, "crops.json")
        with open(path, "w") as fh:
            json.dump(boxes, fh)
        return path

    @unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
    def test_pdf_sizes(self):
        # review_poster.pdf is 720 x 1008 pt. --dpi 150: 1500 x 2100 px raster.
        # full: 2000 x 2000*1008/720 = 2800. Box 0.1-0.6 x 0.1-0.4 = 360 x 302.4 pt -> 750 x 630 px.
        # grid 2x2 tile 360 x 504 pt -> 750 x 1050. print width 50.8 cm = 20 in -> 1500/20 = 75.0 PPI.
        out = os.path.join(self.tmp, "poster_review_x", "crops")
        code, d, _ = run("render_crops.py", "--poster", fx("review_poster.pdf"), "--out-dir", out,
                         "--crops", self.crops_file([BOX]), "--grid", "2x2", "--print-width-cm", "50.8")
        self.assertEqual(code, 0, d)
        self.assertEqual(d["input"]["page_pt"], [720.0, 1008.0])
        self.assertEqual(d["raster_px"], [1500, 2100])
        self.assertEqual(d["full"]["px"], [2000, 2800])
        self.assertEqual(png_size(d["full"]["path"]), [2000, 2800])
        by = {c["name"]: c for c in d["crops"]}
        self.assertEqual(sorted(by), ["intro", "r1c1", "r1c2", "r2c1", "r2c2"])
        self.assertEqual(by["intro"]["px"], [750, 630])
        self.assertEqual(png_size(by["intro"]["path"]), [750, 630])
        for k in ("r1c1", "r1c2", "r2c1", "r2c2"):
            self.assertEqual(by[k]["px"], [750, 1050], k)
        self.assertEqual(by["intro"]["rel"], "crops/intro.png")
        self.assertEqual(by["intro"]["md"], "![intro](crops/intro.png)")
        self.assertEqual(d["effective_ppi"]["at_file_size"], 150.0)
        self.assertAlmostEqual(d["effective_ppi"]["at_print_size"], 75.0, delta=PPI_TOL)

    @unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
    def test_pdf_dpi_and_suffix(self):
        # --dpi 72: raster 720 x 1008 px; px box = (round(72), round(100.8), round(432), round(403.2))
        #  = (72, 101, 432, 403) -> 360 x 302
        code, d, _ = run("render_crops.py", "--poster", fx("review_poster.pdf"), "--out-dir", self.tmp,
                         "--crops", self.crops_file([BOX]), "--dpi", "72", "--suffix", "_before", "--no-full")
        self.assertEqual(code, 0, d)
        self.assertIsNone(d["full"])
        c = d["crops"][0]
        self.assertEqual(c["name"], "intro_before")
        self.assertEqual(c["px"], [360, 302])

    def _png(self, extra_env=None):
        # review_poster.png is 1000 x 1400 px. Box -> (100, 140, 600, 560) -> 500 x 420; grid 2x2 -> 500 x 700.
        # --full-width 500 -> 500 x 700; print width 25.4 cm = 10 in -> 1000/10 = 100 PPI.
        code, d, _ = run("render_crops.py", "--poster", fx("review_poster.png"), "--out-dir", self.tmp,
                         "--crops", self.crops_file([BOX]), "--grid", "2x2", "--full-width", "500",
                         "--print-width-cm", "25.4", env=extra_env or {})
        self.assertEqual(code, 0, d)
        self.assertEqual(d["input"]["px"], [1000, 1400])
        self.assertEqual(d["full"]["px"], [500, 700])
        self.assertEqual(png_size(d["full"]["path"]), [500, 700])
        by = {c["name"]: c for c in d["crops"]}
        self.assertEqual(by["intro"]["px"], [500, 420])
        self.assertEqual(png_size(by["intro"]["path"]), [500, 420])
        self.assertEqual(by["r2c2"]["px"], [500, 700])
        self.assertAlmostEqual(d["effective_ppi"]["at_print_size"], 100.0, delta=PPI_TOL)
        return d

    @unittest.skipUnless(HAS_PIL or HAS_FITZ, "neither Pillow nor PyMuPDF installed")
    def test_png_sizes(self):
        d = self._png()
        self.assertEqual(d["input"]["backend"], "pillow" if HAS_PIL else "pymupdf")

    @unittest.skipUnless(HAS_FITZ, "PyMuPDF not installed")
    def test_png_without_pillow_uses_pymupdf(self):
        blk = os.path.join(self.tmp, "block", "PIL")
        os.makedirs(blk)
        with open(os.path.join(blk, "__init__.py"), "w") as fh:
            fh.write("raise ImportError('blocked for test')\n")
        d = self._png({"PYTHONPATH": os.path.dirname(blk)})
        self.assertEqual(d["input"]["backend"], "pymupdf")

    @unittest.skipUnless(HAS_PIL, "Pillow not installed")
    def test_jpg_input(self):
        from PIL import Image
        jpg = os.path.join(self.tmp, "p.jpg")
        Image.new("RGB", (800, 600), (18, 122, 138)).save(jpg)
        # box (0.25, 0.5, 0.75, 1.0) -> (200, 300, 600, 600) -> 400 x 300; 800 < 2000 -> full not upscaled
        code, d, _ = run("render_crops.py", "--poster", jpg, "--out-dir", self.tmp, "--crops",
                         self.crops_file([{"name": "q", "x0": 0.25, "y0": 0.5, "x1": 0.75, "y1": 1.0}]))
        self.assertEqual(code, 0, d)
        self.assertEqual((d["full"]["px"], d["crops"][0]["px"]), ([800, 600], [400, 300]))

    @unittest.skipUnless(HAS_PIL or HAS_FITZ, "neither Pillow nor PyMuPDF installed")
    def test_invalid_box_reported(self):
        bad = {"name": "bad", "x0": 0.7, "y0": 0.1, "x1": 0.2, "y1": 0.4}
        code, d, _ = run("render_crops.py", "--poster", fx("review_poster.png"), "--out-dir", self.tmp,
                         "--crops", self.crops_file([BOX, bad]))
        self.assertEqual((code, d["ok"]), (1, False))
        self.assertEqual([p["name"] for p in d["problems"]], ["bad"])
        self.assertEqual([c["name"] for c in d["crops"]], ["intro"])          # valid one still written

    def test_canva_url_skipped(self):
        code, d, _ = run("render_crops.py", "--poster", "https://www.canva.com/design/DXXXXXXXXXX/edit",
                         "--out-dir", self.tmp)
        self.assertEqual((code, d["status"]), (2, "skipped"))
        self.assertIn("export-design", d["todo"])

    def test_pptx_skipped_without_soffice(self):
        so = os.environ.get("SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
        if so or os.path.exists("/Applications/LibreOffice.app/Contents/MacOS/soffice"):
            self.skipTest("LibreOffice present: the no-converter path cannot be exercised")
        if not HAS_FITZ:
            self.skipTest("PyMuPDF not installed")
        code, d, _ = run("render_crops.py", "--poster", fx("review_poster.pptx"), "--out-dir", self.tmp,
                         "--no-keynote")
        self.assertEqual((code, d["status"]), (2, "skipped"))
        self.assertIn("PDF", d["todo"])

    @unittest.skipUnless(HAS_FITZ and HAS_PPTX, "PyMuPDF or python-pptx not installed")
    def test_pptx_converted_with_soffice(self):
        if not (shutil.which("soffice") or os.path.exists("/Applications/LibreOffice.app/Contents/MacOS/soffice")):
            self.skipTest("LibreOffice (soffice) not installed")
        from pptx import Presentation
        from pptx.util import Cm
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(25.4), Cm(35.56)              # 10 x 14 in = 720 x 1008 pt
        prs.slides.add_slide(prs.slide_layouts[6])
        src = os.path.join(self.tmp, "p.pptx")
        prs.save(src)
        code, d, _ = run("render_crops.py", "--poster", src, "--out-dir", os.path.join(self.tmp, "c"),
                         "--crops", self.crops_file([BOX]), "--no-keynote")
        self.assertEqual(code, 0, d)
        self.assertEqual(d["input"]["converter"], "soffice")
        self.assertAlmostEqual(d["crops"][0]["px"][0], 750, delta=1)


class TestReviewTemplate(unittest.TestCase):
    def setUp(self):
        with open(os.path.join(TEMPLATES, "REVIEW-template.md"), encoding="utf-8") as fh:
            self.text = fh.read()

    def test_summary_before_first_finding(self):
        summ = self.text.find("## Summary")
        first = re.search(r"(?m)^## N\d+\. ", self.text)
        self.assertGreaterEqual(summ, 0)
        self.assertIsNotNone(first)
        self.assertLess(summ, first.start())
        head = self.text[summ:first.start()]
        for k in ("Verdict", "READY TO PRINT", "FIX FIRST", "must-fix", "should", "nice"):
            self.assertIn(k, head)
        self.assertRegex(head, r"(?m)^1\. ")                                 # ranked top fixes

    def test_preview_tip_near_top(self):
        tip = self.text.find("Markdown preview")
        self.assertGreaterEqual(tip, 0)
        self.assertLess(tip, self.text.find("## Summary"))
        self.assertIn("Cmd+Shift+V", self.text[:tip + 300])

    def test_finding_structure(self):
        self.assertIn("![full poster](crops/full.png)", self.text)
        self.assertRegex(self.text, r"!\[[^\]]*\]\(crops/[\w-]+\.png\)")
        self.assertIn("- [ ] Accept  - [ ] Reject — note:", self.text)
        self.assertIn("_before.png", self.text)
        self.assertIn("_after.png", self.text)
        self.assertRegex(self.text, r"(?m)^- \*\*Severity:\*\*")


# --------------------------------------------------------------------------- missing dependencies
class TestMissingDependencies(Base):
    """Block an optional dependency with a shadow module: the script must exit 2 with status 'skipped'."""

    def blocked_env(self, *mods):
        d = os.path.join(self.tmp, "block")
        os.makedirs(d, exist_ok=True)
        for m in mods:
            os.makedirs(os.path.join(d, m), exist_ok=True)
            with open(os.path.join(d, m, "__init__.py"), "w") as fh:
                fh.write("raise ImportError('blocked for test')\n")
        return {"PYTHONPATH": d}

    def test_check_export_without_pymupdf(self):
        code, d, _ = run("check_export.py", fx("paper.tex"), "--size", "A0", env=self.blocked_env("fitz", "pymupdf"))
        self.assertEqual((code, d["status"]), (2, "skipped"))

    def test_build_pptx_without_python_pptx(self):
        code, d, _ = run("build_pptx.py", fx("pptx_bad_spec.json"), "--out", os.path.join(self.tmp, "x.pptx"),
                         env=self.blocked_env("pptx"))
        self.assertEqual((code, d["status"]), (2, "skipped"))

    def test_export_figs_pdf_without_pymupdf(self):
        code, d, _ = run("export_figs.py", fx("paper.tex").replace("paper.tex", "fig_4x3.pdf"), "--out", self.tmp,
                         env=self.blocked_env("fitz", "pymupdf"))
        self.assertEqual((code, d["status"]), (2, "skipped"))

    def test_render_without_playwright_writes_html(self):
        code, d, _ = run("render_script_text.py", fx("render_ok.json"), "--out", os.path.join(self.tmp, "r.png"),
                         env=self.blocked_env("playwright"))
        self.assertEqual((code, d["status"]), (2, "skipped"))
        self.assertTrue(os.path.exists(d["html"]))

    def test_render_crops_pdf_without_pymupdf(self):
        code, d, _ = run("render_crops.py", "--poster", fx("review_poster.pdf"), "--out-dir", self.tmp,
                         env=self.blocked_env("fitz", "pymupdf"))
        self.assertEqual((code, d["status"]), (2, "skipped"))

    def test_qr_skipped_without_decoders(self):
        if not HAS_FITZ:
            self.skipTest("PyMuPDF not installed")
        code, d, _ = run("check_export.py", fx("clean.pdf"), "--size", "A0", "--qr-url", QR_GOOD,
                         env=self.blocked_env("cv2", "pyzbar"))
        self.assertEqual((code, check(d, "qr")["status"]), (0, "SKIPPED"))


# --------------------------------------------------------------------------- package structure
PORTABLE = {"name", "description", "license", "compatibility", "metadata", "allowed-tools"}


def frontmatter(path):
    with open(path, encoding="utf-8") as fh:
        text = fh.read()
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not m:
        return None, text
    keys = re.findall(r"^([A-Za-z][\w-]*):", m.group(1), re.M)
    return keys, text


class TestStructure(unittest.TestCase):
    def test_skill_frontmatter_and_size(self):
        keys, text = frontmatter(os.path.join(SKILL, "SKILL.md"))
        self.assertIsNotNone(keys)
        self.assertTrue(set(keys) <= PORTABLE, set(keys) - PORTABLE)
        self.assertIn("name", keys)
        self.assertIn("description", keys)
        self.assertRegex(text, r"(?m)^name:\s*academic-poster\s*$")
        self.assertLess(text.count("\n"), 500)

    def test_skill_mentions_every_reference_script_template(self):
        _, text = frontmatter(os.path.join(SKILL, "SKILL.md"))
        for sub in ("references", "scripts", "templates"):
            for f in os.listdir(os.path.join(SKILL, sub)):
                if f.startswith((".", "__")):
                    continue
                self.assertIn(f, text, "%s/%s not mentioned in SKILL.md" % (sub, f))

    def test_skill_named_paths_exist(self):
        _, text = frontmatter(os.path.join(SKILL, "SKILL.md"))
        for sub, name in re.findall(r"\b(references|templates)/([\w.-]+\.(?:md|json))", text):
            self.assertTrue(os.path.exists(os.path.join(SKILL, sub, name)), sub + "/" + name)
        for name in set(re.findall(r"`([a-z_]+\.py)\b", text)):
            self.assertTrue(os.path.exists(os.path.join(SCRIPTS, name)), name)

    def test_agents_frontmatter(self):
        adir = os.path.join(ROOT, "agents")
        files = [f for f in os.listdir(adir) if f.endswith(".md")]
        self.assertEqual(len(files), 4)
        for f in files:
            keys, text = frontmatter(os.path.join(adir, f))
            self.assertIsNotNone(keys, f)
            self.assertIn("name", keys, f)
            self.assertIn("description", keys, f)
            self.assertRegex(text, r"(?m)^name:\s*%s\s*$" % re.escape(f[:-3]))

    def test_readme_install_routes(self):
        with open(os.path.join(ROOT, "README.md"), encoding="utf-8") as fh:
            text = fh.read()
        for route in ("Option A", "Option B", "Option C"):
            self.assertIn(route, text)

    def test_plugin_manifest(self):
        with open(os.path.join(ROOT, ".claude-plugin", "plugin.json"), encoding="utf-8") as fh:
            self.assertEqual(json.load(fh)["name"], "academic-poster")

    def test_templates_are_valid_json(self):
        for f in os.listdir(TEMPLATES):
            if f.endswith(".json"):
                with open(os.path.join(TEMPLATES, f), encoding="utf-8") as fh:
                    json.load(fh)

    def test_python38_syntax(self):
        """No 3.9+ features: builtin generics or X | Y in annotations, match, removeprefix/suffix."""
        files = [os.path.join(SCRIPTS, f) for f in os.listdir(SCRIPTS) if f.endswith(".py")]
        files += [os.path.join(HERE, f) for f in os.listdir(HERE) if f.endswith(".py")]
        for path in files:
            with open(path, encoding="utf-8") as fh:
                src = fh.read()
            tree = ast.parse(src, path)            # under 3.8 this rejects 3.9+ grammar outright
            for node in ast.walk(tree):
                ann = []
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    ann = [a.annotation for a in node.args.args + node.args.kwonlyargs if a.annotation]
                    ann += [node.returns] if node.returns else []
                elif isinstance(node, ast.AnnAssign):
                    ann = [node.annotation]
                for a in ann:
                    for sub in ast.walk(a):
                        if isinstance(sub, ast.Subscript) and isinstance(sub.value, ast.Name):
                            self.assertNotIn(sub.value.id, ("list", "dict", "tuple", "set", "type"), path)
                        self.assertFalse(isinstance(sub, ast.BinOp) and isinstance(sub.op, ast.BitOr), path)
            self.assertNotRegex(src, r"\.remove(prefix|suffix)\(", path)
            self.assertNotRegex(src, r"(?m)^\s*match .+:\s*$", path)


if __name__ == "__main__":
    if not os.path.exists(os.path.join(FX, "manifest.json")):
        sys.exit("fixtures missing: run python3 tests/make_fixtures.py first")
    unittest.main(verbosity=2)
