#!/usr/bin/env python3
"""PPTX route: build an editable poster .pptx from a plan-spec JSON (no Canva needed).

The plan spec is the machine-readable form of a PLAN (see templates/plan-spec.example.json):
  {"page": {"width_cm": 84.1, "height_cm": 118.9, "background": "#ffffff"},
   "units": "cm" | "px",            # px = design px at 96 per inch (Canva custom-size px)
   "min_pt": 24, "min_ppi": 150,
   "fonts":  {"heading": "Arial", "body": "Arial"},
   "colors": {"ink": "#222831"},   # named colours usable anywhere a colour is expected
   "elements": [ ... in Z-order, first = bottom ... ]}
Element types:
  rect | rounded_rect | oval : x, y, w, h, fill, stroke, stroke_pt, radius (0..0.5, rounded_rect)
  text  : x, y, w, h, and ONE of text | slot (key in --wording) | paragraphs
          (a paragraph is a string or a list of [text, {size_pt, bold, italic, color, font}] runs);
          size_pt (or size_px), color, bold, italic, font ("heading" | "body" | a font name),
          align left|center|right, anchor top|middle|bottom, line_spacing
  image : x, y, w, [h], path (relative to the spec file or --assets) | media (key in --media);
          fit "contain" (default; centred in w x h) or "width" (h from aspect ratio)
Checks while building (reported, file still written): text below min_pt, image PPI below min_ppi,
element outside the page, slot/media/file not found.
Output: the .pptx, and JSON on stdout {"ok", "out", "page_cm", "elements", "problems": [...]}.
  Exit 0 = no problems, 1 = problems found, 2 = python-pptx missing (pip install python-pptx).
Import into Canva (optional): upload the .pptx to Canva, which opens PowerPoint files as editable designs
(menu wording varies by Canva version; fonts may be substituted, so re-check text sizes after import).
Python >= 3.8.
"""
import argparse
import json
import os
import struct
import sys

PX_PER_CM = 96 / 2.54


def raster_size(path):
    with open(path, "rb") as fh:
        head = fh.read(26)
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            return struct.unpack(">II", head[16:24])
        fh.seek(0)
        if fh.read(2) != b"\xff\xd8":
            raise ValueError("only PNG/JPEG images are supported in .pptx: %s" % path)
        while True:
            marker = fh.read(2)
            if len(marker) < 2 or marker[0] != 0xFF:
                raise ValueError("bad JPEG: %s" % path)
            if marker[1] in (0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF):
                fh.read(3)
                h, w = struct.unpack(">HH", fh.read(4))
                return w, h
            seg = struct.unpack(">H", fh.read(2))[0]
            fh.seek(seg - 2, 1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="plan-spec JSON")
    ap.add_argument("--out", required=True, help="output .pptx")
    ap.add_argument("--wording", help="wording JSON; text elements may use \"slot\": key")
    ap.add_argument("--media", help="media JSON {key: {\"file\": path, ...}}; images may use \"media\": key")
    ap.add_argument("--assets", help="folder that image paths are relative to (default: spec folder)")
    ap.add_argument("--check-only", action="store_true", help="run the checks, do not write the .pptx")
    a = ap.parse_args()

    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Cm, Pt
    except ImportError:
        print(json.dumps({"ok": None, "status": "skipped",
                          "reason": "dependency missing: python-pptx (pip install python-pptx)"}))
        sys.exit(2)

    with open(a.spec, encoding="utf-8") as fh:
        spec = json.load(fh)
    base = a.assets or os.path.dirname(os.path.abspath(a.spec))
    slots = {}
    if a.wording:
        with open(a.wording, encoding="utf-8") as fh:
            slots = json.load(fh).get("slots", {})
    media = {}
    if a.media:
        with open(a.media, encoding="utf-8") as fh:
            media = json.load(fh)
        media = media.get("media", media)

    page = spec["page"]
    W, H = float(page["width_cm"]), float(page["height_cm"])
    unit = 1.0 / PX_PER_CM if spec.get("units", "cm") == "px" else 1.0
    min_pt, min_ppi = float(spec.get("min_pt", 24)), float(spec.get("min_ppi", 150))
    colors, fonts = spec.get("colors", {}), spec.get("fonts", {})
    problems = []

    def col(c):
        c = colors.get(c, c) if c else c
        return RGBColor.from_string(c.lstrip("#").upper()) if c else None

    def font_name(f):
        return fonts.get(f or "body", f or "Arial")

    def size_pt(d, default=None):
        if "size_pt" in d:
            return float(d["size_pt"])
        if "size_px" in d:
            return float(d["size_px"]) * 0.75
        return default

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(W), Cm(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if page.get("background"):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = col(page["background"])

    shapes = {"rect": MSO_SHAPE.RECTANGLE, "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    anchors = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

    for i, e in enumerate(spec.get("elements", [])):
        t = e.get("type")
        tag = e.get("id", "%s#%d" % (t, i))
        x, y = float(e.get("x", 0)) * unit, float(e.get("y", 0)) * unit
        w, h = float(e.get("w", 0)) * unit, float(e.get("h", 0)) * unit
        if t in shapes:
            sp = slide.shapes.add_shape(shapes[t], Cm(x), Cm(y), Cm(w), Cm(h))
            sp.shadow.inherit = False
            if e.get("fill"):
                sp.fill.solid()
                sp.fill.fore_color.rgb = col(e["fill"])
            else:
                sp.fill.background()
            if e.get("stroke"):
                sp.line.color.rgb = col(e["stroke"])
                sp.line.width = Pt(float(e.get("stroke_pt", 2)))
            else:
                sp.line.fill.background()
            if t == "rounded_rect":
                sp.adjustments[0] = float(e.get("radius", 0.06))
        elif t == "text":
            if "slot" in e:
                if e["slot"] not in slots:
                    problems.append({"element": tag, "problem": "slot not found: %s" % e["slot"]})
                    continue
                sl = slots[e["slot"]]
                txt = (sl.get("short") or {}).get("text") if e.get("use") == "short" else sl.get("text")
                paras = (txt or "").split("\n")
            elif "paragraphs" in e:
                paras = e["paragraphs"]
            else:
                paras = str(e.get("text", "")).split("\n")
            tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.auto_size = None
            tf.vertical_anchor = anchors.get(e.get("anchor", "top"), MSO_ANCHOR.TOP)
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            base_pt = size_pt(e, 28.0)
            for j, para in enumerate(paras):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.alignment = aligns.get(e.get("align", "left"), PP_ALIGN.LEFT)
                p.line_spacing = float(e.get("line_spacing", 1.1))
                runs = [[para, {}]] if isinstance(para, str) else para
                for text, o in runs:
                    r = p.add_run()
                    r.text = text
                    pt = size_pt(o, base_pt)
                    r.font.size = Pt(pt)
                    r.font.bold = bool(o.get("bold", e.get("bold", False)))
                    r.font.italic = bool(o.get("italic", e.get("italic", False)))
                    r.font.name = font_name(o.get("font", e.get("font", "body")))
                    c = o.get("color", e.get("color", "#000000"))
                    r.font.color.rgb = col(c)
                    if text.strip() and pt < min_pt - 0.05:
                        problems.append({"element": tag, "problem": "text %.1f pt < %.0f pt" % (pt, min_pt),
                                         "text": text[:50]})
        elif t == "image":
            path = e.get("path")
            if "media" in e:
                m = media.get(e["media"])
                path = (m or {}).get("file") or (m or {}).get("path") or e["media"]
            if not path:
                problems.append({"element": tag, "problem": "image without path/media"})
                continue
            full = path if os.path.isabs(path) else os.path.join(base, path)
            if not os.path.exists(full):
                problems.append({"element": tag, "problem": "file not found: %s" % full})
                continue
            iw, ih = raster_size(full)
            if e.get("fit") == "width" or not h:
                pw, ph = w, w * ih / float(iw)
                px, py = x, y
            else:
                s = min(w / iw, h / ih)
                pw, ph = iw * s, ih * s
                px, py = x + (w - pw) / 2, y + (h - ph) / 2
            slide.shapes.add_picture(full, Cm(px), Cm(py), Cm(pw), Cm(ph))
            ppi = iw / (pw / 2.54)
            if ppi < min_ppi:
                problems.append({"element": tag, "problem": "image %.0f PPI < %.0f" % (ppi, min_ppi),
                                 "file": path, "px": [iw, ih], "placed_cm": [round(pw, 2), round(ph, 2)]})
            x, y, w, h = px, py, pw, ph
        else:
            problems.append({"element": tag, "problem": "unknown type %r" % t})
            continue
        if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
            problems.append({"element": tag, "problem": "outside the page",
                             "box_cm": [round(x, 2), round(y, 2), round(w, 2), round(h, 2)]})

    if not a.check_only:
        os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
        prs.save(a.out)
    print(json.dumps({"ok": not problems, "out": None if a.check_only else a.out, "page_cm": [W, H],
                      "elements": len(spec.get("elements", [])), "problems": problems},
                     ensure_ascii=False, indent=1))
    sys.exit(0 if not problems else 1)


if __name__ == "__main__":
    main()
